#!/usr/bin/env python3
"""
Post-run verification of the published artefacts.

Counting chart objects through openpyxl proves nothing: a chart can exist as an
object and still draw nothing, because Excel silently refuses to plot a series
whose source cells sit in a hidden column. The only honest check is to unzip the
workbook and read the raw XML that Excel itself will read. Likewise a .hyper
file is only proven by reopening it and running SQL against it.

Checks
------
1. Workbook charts   -- every chart has >= 1 <c:ser>, and plotVisOnly = 0.
2. Workbook links    -- internal jumps carry location=, not an external r:id.
3. Workbook formulas -- the staged policy block matches the pipeline's own CSV,
                        so the live What-If model is reading real numbers.
4. Hyper extract     -- reopened, every table queried, row counts reconciled
                        against the CSVs written beside it.
5. Run manifest      -- headline claims agree with the published CSVs.

Exit code is non-zero if any check fails, so CI can gate on it.

Usage::

    python scripts/verify_outputs.py
"""

from __future__ import annotations

import csv
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Dict, List, Tuple
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "outputs" / "SOP_Command_Center.xlsx"
TABLEAU = ROOT / "outputs" / "tableau"
HYPER = TABLEAU / "SOP_Analytics.hyper"
MANIFEST = ROOT / "outputs" / "run_manifest.json"
PROCESSED = ROOT / "data" / "processed"

C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
S_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

failures: List[str] = []
notes: List[str] = []


def check(condition: bool, ok: str, bad: str) -> None:
    if condition:
        print(f"  PASS  {ok}")
    else:
        print(f"  FAIL  {bad}")
        failures.append(bad)


# --------------------------------------------------------------------------- #
# 1 + 2 + 3. Workbook
# --------------------------------------------------------------------------- #
def verify_workbook() -> None:
    print("\nWORKBOOK  (raw XML inside the .xlsx)")
    if not WORKBOOK.exists():
        check(False, "", f"{WORKBOOK.name} not found — run the pipeline first")
        return

    with zipfile.ZipFile(WORKBOOK) as zf:
        names = zf.namelist()

        # -- charts ------------------------------------------------------- #
        charts = [n for n in names if re.match(r"xl/charts/chart\d+\.xml$", n)]
        empty: List[str] = []
        vis_only: List[str] = []
        total_series = 0

        for name in charts:
            root = ET.fromstring(zf.read(name))
            series = root.iter(f"{C_NS}ser")
            count = sum(1 for _ in series)
            total_series += count
            if count == 0:
                empty.append(name)

            # plotVisOnly must be 0, else hidden staging columns blank the chart
            flags = [e.get("val") for e in root.iter(f"{C_NS}plotVisOnly")]
            if any(v not in ("0", "false") for v in flags):
                vis_only.append(name)

        check(bool(charts), f"{len(charts)} charts embedded",
              "no chart XML parts found in the workbook")
        check(not empty, f"every chart has >=1 <c:ser> ({total_series} series total)",
              f"charts with zero series: {empty}")
        check(not vis_only, "plotVisOnly=0 on every chart (hidden staging still plots)",
              f"charts that would render blank from hidden cells: {vis_only}")

        # -- hyperlinks ---------------------------------------------------- #
        internal = external = 0
        for name in [n for n in names if n.startswith("xl/worksheets/sheet")]:
            root = ET.fromstring(zf.read(name))
            for link in root.iter(f"{S_NS}hyperlink"):
                if link.get("location"):
                    internal += 1
                else:
                    external += 1
        check(internal > 0, f"{internal} internal jumps carry location=",
              "no internal hyperlinks found — back-links would not navigate")
        check(external == 0, "no accidental external link relationships",
              f"{external} hyperlink(s) written as external r:id — Excel ignores these")

        # -- staged policy block vs the pipeline's own CSV ----------------- #
        sheets = [n for n in names if n.startswith("xl/worksheets/sheet")]
        notes.append(f"{len(sheets)} worksheet parts (includes hidden staging)")

    verify_policy_block()
    verify_dashboard_panel()


def verify_dashboard_panel() -> None:
    """
    Reconcile the live control panel's staged block to the run manifest.

    At (All)/(All)/(All) the panel's SUMPRODUCT terms reduce to a plain total
    over the staged per-SKU block, so summing that block in Python is an exact
    check on what Excel will display — without needing Excel to recalculate.
    """
    if not MANIFEST.exists():
        return
    from openpyxl import load_workbook

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb["1 Dashboard"]

    header_row = 1
    columns: Dict[str, int] = {}
    for col in range(96, 120):
        label = ws.cell(row=header_row, column=col).value
        if isinstance(label, str) and label:
            columns[label] = col
    if "SKU_ID" not in columns:
        check(False, "", "dashboard panel staging block not found")
        return

    def column_values(name: str) -> List[float]:
        col = columns[name]
        out: List[float] = []
        row = header_row + 1
        while ws.cell(row=row, column=columns["SKU_ID"]).value is not None:
            value = ws.cell(row=row, column=col).value
            out.append(value)
            row += 1
        return out

    skus = column_values("SKU_ID")
    revenue = sum(float(v or 0) for v in column_values("Revenue_At_Risk"))
    consensus = sum(float(v or 0) for v in column_values("Consensus_Value"))
    red = sum(1 for v in column_values("Risk_Band") if v == "Red")
    expedite = sum(int(v or 0) for v in column_values("Expedite"))

    supply = manifest.get("supply", {})
    check(len(skus) == supply.get("skus_planned"),
          f"panel stages every planned SKU ({len(skus)})",
          f"panel stages {len(skus)} SKUs vs {supply.get('skus_planned')} planned")

    def close(a: float, b: float, tol: float = 1.0) -> bool:
        return abs(a - b) <= tol

    check(close(revenue, supply.get("revenue_at_risk", 0), 1.0),
          f"panel revenue at risk reconciles: Rs {revenue/1e7:.2f} Cr",
          f"panel revenue {revenue:,.0f} vs manifest "
          f"{supply.get('revenue_at_risk', 0):,.0f}")

    check(close(consensus, manifest.get("consensus", {}).get("consensus_value", 0), 1.0),
          f"panel consensus value reconciles: Rs {consensus/1e7:.2f} Cr",
          f"panel consensus {consensus:,.0f} vs manifest "
          f"{manifest.get('consensus', {}).get('consensus_value', 0):,.0f}")

    check(red == supply.get("risk_distribution", {}).get("Red"),
          f"panel Red count reconciles: {red}",
          f"panel Red {red} vs manifest "
          f"{supply.get('risk_distribution', {}).get('Red')}")

    check(expedite == supply.get("expedite_required"),
          f"panel expedite count reconciles: {expedite}",
          f"panel expedite {expedite} vs manifest {supply.get('expedite_required')}")


def verify_policy_block() -> None:
    """The What-If model reads a staged block; prove it holds pipeline numbers."""
    policy_csv = PROCESSED / "inventory_policy.csv"
    if not policy_csv.exists():
        notes.append("inventory_policy.csv absent; skipped formula input check")
        return

    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK, data_only=False)
    if "_policy_data" not in wb.sheetnames:
        check(False, "", "_policy_data staging sheet missing — What-If has no source")
        return

    ws = wb["_policy_data"]
    header = [c.value for c in ws[1]]
    staged: Dict[str, Dict[str, object]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        staged[str(record.get("SKU_ID"))] = record

    with policy_csv.open(newline="", encoding="utf-8") as fh:
        source = {r["SKU_ID"]: r for r in csv.DictReader(fh)}

    check(len(staged) == len(source),
          f"staged policy rows reconcile to inventory_policy.csv ({len(staged)} SKUs)",
          f"staged {len(staged)} rows vs {len(source)} in inventory_policy.csv")

    mismatches = 0
    compared = 0
    for sku, record in list(staged.items())[:50]:
        truth = source.get(sku)
        if not truth:
            mismatches += 1
            continue
        for field in ("Safety_Stock_Units", "Reorder_Point_Units", "MOQ"):
            if field in record and field in truth and record[field] is not None:
                compared += 1
                if abs(float(record[field]) - float(truth[field])) > 0.01:
                    mismatches += 1

    check(mismatches == 0,
          f"{compared} staged values match the pipeline exactly (50-SKU sample)",
          f"{mismatches} staged value(s) diverge from inventory_policy.csv")

    # The What-If sheet must carry live formulas, not pasted values.
    model = wb["7 What-If Model"]
    formulas = sum(
        1 for row in model.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    check(formulas >= 10, f"{formulas} live formulas on the What-If sheet",
          f"only {formulas} formulas — the model is not live")

    # Only formula cells count. The sheet also carries prose explaining why
    # NORMSINV is preferred, and that prose legitimately names NORM.S.INV.
    legacy = sum(
        1 for row in model.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.startswith("=")
        and "NORM.S.INV" in cell.value
    )
    check(legacy == 0, "no NORM.S.INV (would render #NAME? from a non-Excel writer)",
          f"{legacy} formula(s) use NORM.S.INV instead of NORMSINV")


# --------------------------------------------------------------------------- #
# 4. Hyper extract
# --------------------------------------------------------------------------- #
def verify_hyper() -> None:
    print("\nTABLEAU EXTRACT  (reopened and queried)")
    if not HYPER.exists():
        check(False, "", f"{HYPER.name} not found")
        return

    try:
        from tableauhyperapi import Connection, HyperProcess, Telemetry
    except ImportError:
        notes.append("tableauhyperapi not installed; .hyper query skipped")
        return

    csv_rows: Dict[str, int] = {}
    for path in sorted((TABLEAU / "data").glob("*.csv")):
        with path.open(newline="", encoding="utf-8") as fh:
            csv_rows[path.stem.lower()] = sum(1 for _ in fh) - 1

    with HyperProcess(telemetry=Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU) as hp:
        with Connection(endpoint=hp.endpoint, database=HYPER) as conn:
            schemas = conn.catalog.get_schema_names()
            tables = [t for s in schemas for t in conn.catalog.get_table_names(s)]
            check(bool(tables), f"{len(tables)} tables in the extract",
                  "extract opened but contains no tables")

            total = 0
            mismatched: List[str] = []
            for table in tables:
                count = conn.execute_scalar_query(
                    f"SELECT COUNT(*) FROM {table}")
                total += count
                stem = str(table.name).strip('"').lower()
                if stem in csv_rows and csv_rows[stem] != count:
                    mismatched.append(
                        f"{stem}: hyper {count} vs csv {csv_rows[stem]}")

            check(not mismatched,
                  f"every table reconciles to its CSV ({total:,} rows queried)",
                  f"row-count mismatches: {mismatched}")

            # A real query, not just a count -- proves the data is usable.
            flat = [t for t in tables if "flat_sku_week" in str(t.name).lower()]
            if flat:
                sample = conn.execute_list_query(
                    f"SELECT COUNT(DISTINCT \"SKU_ID\") FROM {flat[0]}")
                check(sample[0][0] > 0,
                      f"flat_sku_week queryable: {sample[0][0]} distinct SKUs",
                      "flat_sku_week returned no SKUs")


# --------------------------------------------------------------------------- #
# 5. Manifest reconciliation
# --------------------------------------------------------------------------- #
def verify_manifest() -> None:
    print("\nRUN MANIFEST  (claims vs published CSVs)")
    if not MANIFEST.exists():
        check(False, "", "run_manifest.json not found")
        return

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    forecast = manifest.get("forecast", {})

    accuracy_csv = PROCESSED / "forecast_accuracy_hierarchy.csv"
    if accuracy_csv.exists():
        with accuracy_csv.open(newline="", encoding="utf-8") as fh:
            levels = list(csv.DictReader(fh))
        check(len(levels) >= 6,
              f"accuracy reported at {len(levels)} aggregation levels",
              f"only {len(levels)} aggregation levels — spec requires 6")
        has_fva = any(
            k for k in (levels[0].keys() if levels else []) if "FVA" in k.upper())
        check(has_fva, "forecast value added reported alongside every level",
              "no FVA column in the accuracy table")

    coverage = forecast.get("interval_coverage", {})
    check(bool(coverage),
          f"interval coverage measured: 80% -> {coverage.get('coverage_80', 0):.1f}%, "
          f"95% -> {coverage.get('coverage_95', 0):.1f}%",
          "no measured interval coverage in the manifest")

    mix = forecast.get("champion_mix", {})
    check(len(mix) >= 3,
          f"champion mix spans {len(mix)} model families: {mix}",
          f"champion mix collapsed to {len(mix)} model(s)")


def main() -> int:
    print("=" * 74)
    print("ARTEFACT VERIFICATION")
    print("=" * 74)

    verify_workbook()
    verify_hyper()
    verify_manifest()

    print("\n" + "=" * 74)
    for note in notes:
        print(f"  note  {note}")
    if failures:
        print(f"VERIFICATION FAILED — {len(failures)} check(s) did not pass")
        for failure in failures:
            print(f"    - {failure}")
        return 1
    print("VERIFICATION PASSED — all checks green")
    return 0


if __name__ == "__main__":
    sys.exit(main())
